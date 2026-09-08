from datetime import datetime
from copy import deepcopy
from io import BytesIO
from html import escape
from urllib.parse import quote
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from pydantic import Field
from sqlalchemy import String, JSON, Integer, Boolean, DateTime, ForeignKey, select, update
from sqlalchemy.orm import Mapped, mapped_column, Session
from .database import Base
from .dependencies import session
from .models import identifier, utcnow, iso, AuditEvent
from .profile_schemas import StrictModel
from .versions import ResumeVersion

class CoverLetter(Base):
    __tablename__='cover_letters'
    id:Mapped[str]=mapped_column(String(36),primary_key=True,default=identifier)
    name:Mapped[str]=mapped_column(String(150))
    document:Mapped[dict]=mapped_column(JSON)
    resume_version_id:Mapped[str|None]=mapped_column(ForeignKey('resume_versions.id',ondelete='SET NULL'),nullable=True)
    archived:Mapped[bool]=mapped_column(Boolean,default=False)
    revision:Mapped[int]=mapped_column(Integer,default=1)
    updated_at:Mapped[datetime]=mapped_column(DateTime(timezone=True),default=utcnow)

class CoverVersion(Base):
    __tablename__='cover_letter_versions'
    id:Mapped[str]=mapped_column(String(36),primary_key=True,default=identifier)
    cover_id:Mapped[str]=mapped_column(ForeignKey('cover_letters.id',ondelete='CASCADE'),index=True)
    snapshot:Mapped[dict]=mapped_column(JSON)
    note:Mapped[str]=mapped_column(String(1000),default='')
    created_at:Mapped[datetime]=mapped_column(DateTime(timezone=True),default=utcnow)

class LetterDocument(StrictModel):
    header:str=Field(default='',max_length=2000)
    greeting:str=Field(default='Dear Hiring Team,',max_length=500)
    opening:str=Field(default='',max_length=10000)
    experience:str=Field(default='',max_length=10000)
    company_fit:str=Field(default='',max_length=10000)
    closing:str=Field(default='Sincerely,',max_length=2000)
    job_description:str=Field(default='',max_length=30000)

class LetterInput(StrictModel):
    name:str=Field(min_length=1,max_length=150)
    document:LetterDocument=Field(default_factory=LetterDocument)
    resume_version_id:str|None=None
    archived:bool=False
    revision:int=Field(default=1,ge=1)

class Note(StrictModel):
    note:str=Field(default='',max_length=1000)

def letter_json(row):return {'id':row.id,'name':row.name,'document':row.document,'resume_version_id':row.resume_version_id,'archived':row.archived,'revision':row.revision,'updated_at':iso(row.updated_at)}
def require(db,id):
    row=db.get(CoverLetter,id)
    if not row:raise HTTPException(404,'Cover letter not found')
    return row

def validate_link(db,p):
    if p.resume_version_id and not db.get(ResumeVersion,p.resume_version_id):raise HTTPException(422,'Resume version not found')

router=APIRouter(prefix='/api/cover-letters',tags=['Cover letters'])
@router.get('')
def listing(db:Session=Depends(session)):return [letter_json(r) for r in db.scalars(select(CoverLetter).order_by(CoverLetter.updated_at.desc()))]
@router.post('',status_code=201)
def create(p:LetterInput,db:Session=Depends(session)):
    validate_link(db,p);row=CoverLetter(name=p.name,document=p.document.model_dump(),resume_version_id=p.resume_version_id);db.add(row);db.flush();db.add(AuditEvent(action='Cover Letter Created',entity_id=row.id));db.commit();return letter_json(row)
@router.put('/{id}')
def edit(id:str,p:LetterInput,db:Session=Depends(session)):
    require(db,id);validate_link(db,p)
    result=db.execute(update(CoverLetter).where(CoverLetter.id==id,CoverLetter.revision==p.revision).values(name=p.name,document=p.document.model_dump(),resume_version_id=p.resume_version_id,archived=p.archived,revision=p.revision+1,updated_at=utcnow()))
    if result.rowcount!=1:raise HTTPException(409,'Cover letter changed. Reload before saving.')
    db.commit();return letter_json(require(db,id))
@router.post('/{id}/duplicate',status_code=201)
def duplicate(id:str,db:Session=Depends(session)):
    source=require(db,id);row=CoverLetter(name=source.name[:143]+' (copy)',document=deepcopy(source.document),resume_version_id=source.resume_version_id);db.add(row);db.commit();return letter_json(row)
@router.delete('/{id}',status_code=204)
def remove(id:str,db:Session=Depends(session)):db.delete(require(db,id));db.commit()
@router.get('/{id}/versions')
def versions(id:str,db:Session=Depends(session)):
    require(db,id);return [{'id':v.id,'note':v.note,'snapshot':v.snapshot,'created_at':iso(v.created_at)} for v in db.scalars(select(CoverVersion).where(CoverVersion.cover_id==id).order_by(CoverVersion.created_at.desc()))]
@router.post('/{id}/versions',status_code=201)
def version(id:str,p:Note,db:Session=Depends(session)):
    row=require(db,id);v=CoverVersion(cover_id=id,note=p.note,snapshot=letter_json(row));db.add(v);db.add(AuditEvent(action='Cover Letter Version Created',entity_id=id));db.commit();return {'id':v.id}
@router.post('/{id}/export/{format}')
def export(id:str,format:str,db:Session=Depends(session)):
    row=require(db,id);parts=[row.document.get(k,'') for k in ['header','greeting','opening','experience','company_fit','closing']]
    if format=='docx':
        from docx import Document
        document=Document()
        for part in parts:
            for line in part.split('\n'):document.add_paragraph(line)
        document.core_properties.author='';buffer=BytesIO();document.save(buffer);content=buffer.getvalue();media='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    elif format=='pdf':
        from playwright.sync_api import sync_playwright
        from .exports import export_slots
        with export_slots,sync_playwright() as p:
            try:browser=p.chromium.launch()
            except Exception:browser=p.chromium.launch(channel='chrome')
            try:
                page=browser.new_page();page.route('**/*',lambda route:route.abort());page.set_content('<style>@page{size:A4;margin:22mm}body{font:11pt/1.6 Arial;color:#243c33}p{white-space:pre-wrap;orphans:3;widows:3}</style>'+''.join('<p>'+escape(part)+'</p>' for part in parts));content=page.pdf(format='A4',prefer_css_page_size=True,tagged=True);media='application/pdf'
            finally:browser.close()
    else:raise HTTPException(422,'Choose pdf or docx')
    db.add(AuditEvent(action='Cover Letter Exported',entity_id=id));db.commit()
    return Response(content,media_type=media,headers={'Content-Disposition':f"attachment; filename*=UTF-8''{quote(row.name,safe='')}.{format}",'Cache-Control':'no-store'})
