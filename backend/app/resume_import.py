from __future__ import annotations

import base64
import binascii
import io
import re
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal
from uuid import uuid4

from docx import Document
from fastapi import APIRouter, Depends, HTTPException
from pydantic import Field
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session

from .dependencies import session
from .models import AuditEvent
from .profile import get_profile
from .profile_models import CareerItem
from .profile_schemas import KINDS, ItemInput, PersonalDetails, StrictModel

router = APIRouter(prefix="/api/import", tags=["Resume import"])

MAX_UPLOAD_BYTES = 8 * 1024 * 1024
MAX_TEXT_CHARS = 250_000
MAX_PDF_PAGES = 100

SECTION_ALIASES = {
    "summary": {"summary", "profile", "professional summary", "career summary", "objective"},
    "experience": {
        "experience",
        "work experience",
        "professional experience",
        "employment",
        "employment history",
        "career history",
    },
    "education": {"education", "academic background", "academic history", "qualifications"},
    "skills": {"skills", "technical skills", "core skills", "competencies", "technologies", "tools"},
    "projects": {"projects", "selected projects", "personal projects", "technical projects"},
    "certifications": {"certifications", "certificates", "licenses", "licenses and certifications"},
    "achievements": {"achievements", "awards", "honors", "honours", "accomplishments"},
    "publications": {"publications", "research", "research publications", "papers"},
    "languages": {"languages", "language proficiency"},
    "portfolio": {"portfolio", "selected work"},
}
ALIAS_TO_KIND = {alias: kind for kind, aliases in SECTION_ALIASES.items() for alias in aliases}
DATE_WORD = r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|jul(?:y)?|aug(?:ust)?|sep(?:t(?:ember)?)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)"
DATE_TOKEN = rf"(?:{DATE_WORD}\s+\d{{4}}|\d{{1,2}}[/-]\d{{4}}|\d{{4}}|present|current|now)"
DATE_RANGE_RE = re.compile(
    rf"(?P<start>{DATE_TOKEN})\s*(?:-|–|—|to)\s*(?P<end>{DATE_TOKEN})", re.I
)
EMAIL_RE = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.I)
# Exclude common employment/education year ranges before accepting a phone-like token.
PHONE_RE = re.compile(
    r"(?<!\w)(?!(?:19|20)\d{2}\s*[-–—]\s*(?:19|20)\d{2}(?!\d))"
    r"(?:\+?\d[\d\s().-]{6,}\d)(?!\w)"
)
URL_RE = re.compile(r"https?://[^\s<>\])}]+", re.I)
DOI_RE = re.compile(r"\b10\.\d{4,9}/[-._;()/:A-Z0-9]+\b", re.I)


@dataclass
class ParsedDocument:
    text: str
    format: str
    page_count: int = 1
    image_count: int = 0
    link_count: int = 0
    links: list[str] = field(default_factory=list)
    table_count: int = 0
    table_text_ratio: float = 0.0
    scanned: bool = False
    likely_multi_column: bool = False
    warnings: list[str] = field(default_factory=list)


class UploadPayload(StrictModel):
    filename: str = Field(min_length=1, max_length=255)
    content_base64: str = Field(min_length=1, max_length=MAX_UPLOAD_BYTES * 2)


class ImportSelection(StrictModel):
    temp_id: str = Field(max_length=100)
    kind: str = Field(max_length=30)
    data: dict
    selected: bool = True
    duplicate_id: str | None = Field(default=None, max_length=100)
    duplicate_action: Literal["keep", "replace", "merge", "new"] = "keep"


class ApplyImport(StrictModel):
    filename: str = Field(min_length=1, max_length=255)
    personal: dict[str, str] = Field(default_factory=dict)
    personal_selected: dict[str, bool] = Field(default_factory=dict)
    items: list[ImportSelection] = Field(default_factory=list, max_length=1000)


def _safe_decode(payload: UploadPayload) -> bytes:
    try:
        data = base64.b64decode(payload.content_base64, validate=True)
    except (binascii.Error, ValueError) as exc:
        raise HTTPException(422, "The uploaded file could not be decoded") from exc
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "Resume files are limited to 8 MB")
    return data


def _safe_docx_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 5000 or sum(info.file_size for info in infos) > 32 * 1024 * 1024:
                raise HTTPException(413, "The DOCX expands beyond the safe processing limit")
    except zipfile.BadZipFile as exc:
        raise HTTPException(422, "The DOCX file is corrupted or unsupported") from exc


def _parse_pdf(data: bytes) -> ParsedDocument:
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise HTTPException(422, "The PDF is corrupted or unsupported") from exc
    if reader.is_encrypted:
        try:
            if reader.decrypt("") == 0:
                raise HTTPException(422, "Encrypted PDFs must be unlocked before import")
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, "Encrypted PDFs must be unlocked before import") from exc
    if len(reader.pages) > MAX_PDF_PAGES:
        raise HTTPException(413, f"PDFs are limited to {MAX_PDF_PAGES} pages")

    parts: list[str] = []
    images = 0
    links: set[str] = set()
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            parts.append("")
        try:
            images += len(page.images)
        except Exception:
            pass
        try:
            for ref in page.get("/Annots", []) or []:
                annotation = ref.get_object()
                action = annotation.get("/A")
                if action and action.get("/URI"):
                    links.add(str(action.get("/URI")))
        except Exception:
            pass

    text = "\n".join(parts).strip()[:MAX_TEXT_CHARS]
    links.update(URL_RE.findall(text))
    chars_per_page = len(text) / max(len(reader.pages), 1)
    scanned = len(text) < 80 or (images > 0 and chars_per_page < 120)
    spaced_lines = [line for line in text.splitlines() if re.search(r"\S\s{5,}\S", line)]
    likely_multi = len(spaced_lines) >= max(3, len(text.splitlines()) // 8)
    warnings: list[str] = []
    if scanned:
        warnings.append(
            "This PDF contains very little selectable text and appears image-based or scanned. "
            "CareerCanvas does not run OCR automatically; use a text-based PDF or local OCR before import."
        )
    return ParsedDocument(
        text=text,
        format="PDF",
        page_count=max(len(reader.pages), 1),
        image_count=images,
        link_count=len(links),
        links=sorted(links),
        scanned=scanned,
        likely_multi_column=likely_multi,
        warnings=warnings,
    )


def _parse_docx(data: bytes) -> ParsedDocument:
    _safe_docx_zip(data)
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise HTTPException(422, "The DOCX file is corrupted or unsupported") from exc

    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    table_text: list[str] = []
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                value = cell.text.strip()
                if value:
                    table_text.append(value)
    text = "\n".join(paragraphs + table_text).strip()[:MAX_TEXT_CHARS]
    links = set(URL_RE.findall(text))
    try:
        for relationship in document.part.rels.values():
            if "hyperlink" in relationship.reltype and relationship.target_ref:
                links.add(str(relationship.target_ref))
    except Exception:
        pass
    table_chars = sum(len(value) for value in table_text)
    ratio = table_chars / max(len(text), 1)
    warnings: list[str] = []
    if document.tables and ratio > 0.45:
        warnings.append(
            "A large share of this DOCX is stored in tables. Some applicant tracking systems may read complex tables unpredictably."
        )
    return ParsedDocument(
        text=text,
        format="DOCX",
        page_count=1,
        link_count=len(links),
        links=sorted(links),
        table_count=len(document.tables),
        table_text_ratio=round(ratio, 3),
        likely_multi_column=bool(document.tables) and ratio > 0.45,
        warnings=warnings,
    )


def _parse_txt(data: bytes) -> ParsedDocument:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = data.decode("cp1252")
        except UnicodeDecodeError as exc:
            raise HTTPException(422, "TXT files must use UTF-8 or Windows-1252 text encoding") from exc
    text = text.strip()[:MAX_TEXT_CHARS]
    links = sorted(set(URL_RE.findall(text)))
    return ParsedDocument(text=text, format="TXT", links=links, link_count=len(links))


def parse_document(filename: str, data: bytes) -> ParsedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(data)
    if suffix == ".docx":
        return _parse_docx(data)
    if suffix == ".txt":
        return _parse_txt(data)
    raise HTTPException(415, "Supported resume formats are PDF, DOCX, and TXT")


def _norm(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()


def _heading(line: str) -> str | None:
    candidate = _norm(line.rstrip(":"))
    return ALIAS_TO_KIND.get(candidate) if len(candidate) <= 45 else None


def _clean_lines(text: str) -> list[str]:
    return [
        re.sub(r"^[\s•●▪◦·*-]+", "", re.sub(r"\s+", " ", line)).strip()
        for line in text.replace("\r", "\n").split("\n")
        if line.strip()
    ]


def split_sections(text: str) -> tuple[list[str], dict[str, list[str]], list[str]]:
    preamble: list[str] = []
    sections: dict[str, list[str]] = defaultdict(list)
    detected: list[str] = []
    current: str | None = None
    for line in _clean_lines(text):
        kind = _heading(line)
        if kind:
            current = kind
            if kind not in detected:
                detected.append(kind)
        elif current:
            sections[current].append(line)
        else:
            preamble.append(line)
    return preamble, dict(sections), detected


def extract_personal(lines: list[str], full_text: str) -> dict[str, dict]:
    result: dict[str, dict] = {}
    email = EMAIL_RE.search(full_text)
    phone = PHONE_RE.search(full_text)
    urls = list(dict.fromkeys(URL_RE.findall(full_text)))
    if email:
        result["email"] = {"value": email.group(0), "confidence": 0.99, "source": email.group(0)}
    if phone:
        result["phone"] = {"value": phone.group(0), "confidence": 0.94, "source": phone.group(0)}
    for url in urls:
        key = (
            "linkedin"
            if "linkedin.com" in url.lower()
            else "github"
            if "github.com" in url.lower()
            else "portfolio"
            if any(term in url.lower() for term in ("portfolio", "behance", "dribbble"))
            else "website"
        )
        result.setdefault(key, {"value": url, "confidence": 0.96, "source": url})

    candidates = [
        line
        for line in lines[:12]
        if not EMAIL_RE.search(line)
        and not PHONE_RE.search(line)
        and not URL_RE.search(line)
        and not DATE_RANGE_RE.search(line)
    ]
    if candidates:
        name = candidates[0]
        words = name.split()
        if 2 <= len(words) <= 6 and len(name) <= 100 and sum(c.isalpha() for c in name) >= len(name) * 0.6:
            result["full_name"] = {"value": name, "confidence": 0.78, "source": name}
            candidates = candidates[1:]
    if candidates:
        title = candidates[0]
        if len(title) <= 120 and not re.search(r"\d{4}", title):
            result["professional_title"] = {"value": title, "confidence": 0.68, "source": title}
    for line in lines[:15]:
        if "," in line and len(line) <= 100 and not EMAIL_RE.search(line) and not URL_RE.search(line):
            parts = [part.strip() for part in line.split(",") if part.strip()]
            if len(parts) == 2 and not any(character.isdigit() for character in line):
                result.setdefault("city", {"value": parts[0], "confidence": 0.55, "source": line})
                result.setdefault("country", {"value": parts[1], "confidence": 0.55, "source": line})
                break
    return result


def _date_parts(line: str) -> tuple[str, str, bool]:
    match = DATE_RANGE_RE.search(line)
    if not match:
        return "", "", False
    start, end = match.group("start"), match.group("end")
    current = end.lower() in {"present", "current", "now"}
    return start, "" if current else end, current


def _split_header(line: str) -> tuple[str, str]:
    for separator in (" | ", " — ", " – ", " @ ", " - "):
        if separator in line:
            left, right = [part.strip() for part in line.split(separator, 1)]
            if left and right:
                return left, right
    return line.strip(), ""


def parse_experience(lines: list[str]) -> list[dict]:
    results: list[dict] = []
    date_indices = [index for index, line in enumerate(lines) if DATE_RANGE_RE.search(line)]
    for position, index in enumerate(date_indices):
        date_line = lines[index]
        start, end, current = _date_parts(date_line)
        inline = DATE_RANGE_RE.sub("", date_line).strip(" |–—-")
        header = inline or (lines[index - 1] if index > 0 else "")
        role, company = _split_header(header)
        if not company and index > 1 and lines[index - 2] != header:
            company = lines[index - 2]
        if not role or not company:
            continue
        next_index = date_indices[position + 1] if position + 1 < len(date_indices) else len(lines)
        bullets = [
            {"id": str(uuid4()), "text": line}
            for line in lines[index + 1 : next_index]
            if len(line) > 8 and not DATE_RANGE_RE.search(line)
        ][:12]
        results.append(
            {
                "data": {
                    "company": company[:200],
                    "position": role[:200],
                    "location": "",
                    "start_date": start,
                    "end_date": end,
                    "current": current,
                    "description": "",
                    "bullets": bullets,
                    "technologies": [],
                },
                "confidence": 0.76 if inline else 0.67,
                "source": " | ".join(lines[max(0, index - 2) : min(len(lines), index + 3)])[:1000],
            }
        )
    return results


def parse_education(lines: list[str]) -> list[dict]:
    results: list[dict] = []
    date_indices = [index for index, line in enumerate(lines) if DATE_RANGE_RE.search(line)]
    if not date_indices and lines:
        date_indices = [len(lines)]
    for position, index in enumerate(date_indices):
        end_at = date_indices[position + 1] if position + 1 < len(date_indices) else len(lines)
        before = lines[max(0, index - 2) : index] if index < len(lines) else lines[:2]
        if len(before) < 2:
            continue
        institution, degree = before[-2], before[-1]
        start, end, _ = _date_parts(lines[index]) if index < len(lines) else ("", "", False)
        rest = lines[index + 1 : end_at] if index < len(lines) else lines[2:]
        gpa = next(
            (
                match.group(1)
                for line in rest
                if (match := re.search(r"GPA\s*[: ]\s*([\d.]+(?:\s*/\s*[\d.]+)?)", line, re.I))
            ),
            "",
        )
        results.append(
            {
                "data": {
                    "institution": institution[:200],
                    "degree": degree[:200],
                    "field": "",
                    "start_date": start,
                    "end_date": end,
                    "gpa": gpa,
                    "location": "",
                    "description": "",
                    "courses": [],
                    "honors": [],
                },
                "confidence": 0.66,
                "source": " | ".join(before + ([lines[index]] if index < len(lines) else []))[:1000],
            }
        )
    return results


def parse_skills(lines: list[str]) -> list[dict]:
    results: list[dict] = []
    seen: set[str] = set()
    for line in lines:
        source = line
        line = re.sub(r"^[A-Za-z /&+#-]{2,30}:\s*", "", line)
        for value in re.split(r"[,;|•·]", line):
            value = value.strip(" .")
            key = _norm(value)
            if not value or len(value) > 80 or key in seen or len(key) < 2 or len(value.split()) > 6:
                continue
            seen.add(key)
            results.append(
                {
                    "data": {
                        "name": value[:100],
                        "category": "Other",
                        "level": "",
                        "target_level": "",
                        "target_date": "",
                        "learning_notes": "",
                    },
                    "confidence": 0.78,
                    "source": source[:500],
                }
            )
    return results[:100]


def parse_projects(lines: list[str]) -> list[dict]:
    results: list[dict] = []
    current: dict | None = None
    for line in lines:
        if URL_RE.fullmatch(line):
            if current:
                current["data"]["github" if "github.com" in line.lower() else "url"] = line
            continue
        is_title = (
            len(line) <= 100
            and not line.endswith(".")
            and not re.match(r"^(built|developed|created|implemented|designed|led|used)\b", line, re.I)
        )
        if is_title:
            if current:
                results.append(current)
            current = {
                "data": {
                    "name": line[:200],
                    "role": "",
                    "description": "",
                    "start_date": "",
                    "end_date": "",
                    "url": "",
                    "github": "",
                    "technologies": [],
                    "bullets": [],
                },
                "confidence": 0.61,
                "source": line[:500],
            }
        elif current and len(line) > 5:
            current["data"]["bullets"].append({"id": str(uuid4()), "text": line[:3000]})
            current["source"] = (current["source"] + " | " + line)[:1000]
    if current:
        results.append(current)
    return results[:50]


def parse_portfolio(lines: list[str]) -> list[dict]:
    results: list[dict] = []
    pending_title = ""
    for line in lines:
        urls = URL_RE.findall(line)
        if urls:
            for url in urls:
                title = line.replace(url, "").strip(" |–—-:") or pending_title or url
                results.append(
                    {
                        "data": {
                            "title": title[:200],
                            "description": "",
                            "kind": "Project",
                            "thumbnail": "",
                            "skills": [],
                            "url": "" if "github.com" in url.lower() else url,
                            "github": url if "github.com" in url.lower() else "",
                            "date": "",
                            "tags": [],
                        },
                        "confidence": 0.7,
                        "source": line[:500],
                    }
                )
            pending_title = ""
        elif len(line) <= 120:
            pending_title = line
    return results[:50]


def _simple_line_items(lines: list[str], kind: str) -> list[dict]:
    results: list[dict] = []
    for line in lines:
        if len(line) < 2:
            continue
        if kind == "certifications":
            left, right = _split_header(line)
            data = {
                "name": left[:200],
                "issuer": right[:200],
                "date": "",
                "expiry": "",
                "credential_id": "",
                "credential_url": "",
            }
        elif kind == "achievements":
            data = {
                "statement": line[:3000],
                "company": "",
                "project": "",
                "skills": [],
                "metric": "",
                "category": "Impact",
                "tags": [],
            }
        elif kind == "publications":
            doi = DOI_RE.search(line)
            data = {
                "title": line[:500],
                "authors": "",
                "venue": "",
                "year": next(iter(re.findall(r"\b(?:19|20)\d{2}\b", line)), ""),
                "doi": doi.group(0) if doi else "",
                "url": next(iter(URL_RE.findall(line)), ""),
                "citation": line[:2000],
                "description": "",
            }
        elif kind == "languages":
            left, right = _split_header(line.replace(":", " | ", 1) if ":" in line else line)
            data = {"language": left[:100], "proficiency": right[:100]}
        else:
            continue
        results.append({"data": data, "confidence": 0.62, "source": line[:500]})
    return results[:100]


def extract_items(sections: dict[str, list[str]]) -> list[dict]:
    items: list[dict] = []
    parsers = {
        "experience": parse_experience,
        "education": parse_education,
        "skills": parse_skills,
        "projects": parse_projects,
        "portfolio": parse_portfolio,
    }
    for kind, lines in sections.items():
        if kind == "summary":
            continue
        parsed = parsers[kind](lines) if kind in parsers else _simple_line_items(lines, kind)
        for candidate in parsed:
            try:
                validated = ItemInput(kind=kind, data=candidate["data"])
            except Exception:
                continue
            items.append(
                {
                    "temp_id": str(uuid4()),
                    "kind": kind,
                    "data": validated.data,
                    "confidence": candidate["confidence"],
                    "source": candidate["source"],
                }
            )
    return items


def _title(item: CareerItem) -> str:
    for key in ("name", "position", "institution", "statement", "title", "language"):
        if item.data.get(key):
            return str(item.data[key])
    return item.kind.title()


def _duplicate_score(kind: str, incoming: dict, existing: CareerItem) -> float:
    data = existing.data
    if kind == "skills":
        return 1.0 if _norm(str(incoming.get("name", ""))) == _norm(str(data.get("name", ""))) else 0.0
    if kind == "experience":
        company = _norm(str(incoming.get("company", ""))) == _norm(str(data.get("company", "")))
        position = _norm(str(incoming.get("position", ""))) == _norm(str(data.get("position", "")))
        return 0.96 if company and position else 0.82 if company else 0.0
    keys = {
        "education": ("institution", "degree"),
        "projects": ("name",),
        "certifications": ("name", "issuer"),
        "publications": ("title",),
        "achievements": ("statement",),
        "languages": ("language",),
    }.get(kind, ())
    if not keys:
        return 0.0
    matches = sum(
        _norm(str(incoming.get(key, ""))) == _norm(str(data.get(key, "")))
        and bool(_norm(str(incoming.get(key, ""))))
        for key in keys
    )
    return matches / len(keys)


def attach_duplicates(items: list[dict], db: Session) -> None:
    profile = get_profile(db)
    existing = db.scalars(select(CareerItem).where(CareerItem.profile_id == profile.id)).all()
    by_kind: dict[str, list[CareerItem]] = defaultdict(list)
    for item in existing:
        by_kind[item.kind].append(item)
    for candidate in items:
        scored = [
            (_duplicate_score(candidate["kind"], candidate["data"], item), item)
            for item in by_kind[candidate["kind"]]
        ]
        score, match = max(scored, default=(0.0, None), key=lambda pair: pair[0])
        candidate["duplicate"] = (
            {"id": match.id, "title": _title(match), "score": round(score, 2)}
            if match is not None and score >= 0.8
            else None
        )


def uploaded_ats(parsed: ParsedDocument, detected: list[str], personal: dict[str, dict]) -> dict:
    checks: list[dict] = []

    def add(key: str, label: str, weight: int, passed: bool, severity: str, message: str) -> None:
        checks.append(
            {
                "key": key,
                "label": label,
                "weight": weight,
                "passed": bool(passed),
                "severity": severity,
                "message": message,
                "earned": weight if passed else 0,
            }
        )

    detected_set = set(detected)
    text_lines = _clean_lines(parsed.text)
    candidate_headings = [
        line for line in text_lines if len(line) <= 45 and (line.isupper() or line.istitle())
    ]
    unusual = [
        line
        for line in candidate_headings
        if _heading(line) is None and len(line.split()) <= 5
    ][:12]
    add("text", "Extractable text available", 10, len(parsed.text.strip()) > 80, "HIGH", "Use a text-based PDF/DOCX so applicant tracking systems can read the resume content.")
    add("headings", "Standard section headings", 10, {"experience", "education", "skills"}.issubset(detected_set), "MEDIUM", "Use recognizable headings such as Experience, Education, Skills, and Projects.")
    add("columns", "Simple reading order", 10, not parsed.likely_multi_column, "MEDIUM", "Complex columns or table-heavy layouts can produce less predictable reading order.")
    add("images", "Critical content is extractable text", 5, not parsed.scanned and len(parsed.text) > 120, "HIGH", "Do not place essential qualifications or contact details only inside images.")
    add("email", "Contact email present", 15, "email" in personal, "HIGH", "Add a valid, visible contact email.")
    add("phone", "Contact phone present", 5, "phone" in personal and len(re.sub(r"\D", "", personal["phone"]["value"])) >= 7, "MEDIUM", "Add a visible phone number when appropriate for your application.")
    add("experience", "Experience section detected", 10, "experience" in detected_set, "MEDIUM", "Use a clear Experience or Professional Experience section when applicable.")
    add("education", "Education section detected", 10, "education" in detected_set, "MEDIUM", "Use a clear Education section.")
    add("skills", "Skills section detected", 10, "skills" in detected_set, "MEDIUM", "Use a clear Skills or Technical Skills section.")
    chars_per_page = len(parsed.text) / max(parsed.page_count, 1)
    add("readability", "Readable text density", 10, chars_per_page >= 250 and not parsed.scanned, "MEDIUM", "Very sparse extracted text can indicate image-based content or a parser-unfriendly layout.")
    add("pages", "Reasonable page count", 5, 1 <= parsed.page_count <= 2, "LOW", "Consider a focused one- or two-page resume. Longer academic CVs may be appropriate.")
    warnings = list(parsed.warnings)
    if unusual:
        warnings.append("Possible non-standard headings: " + ", ".join(unusual))
    return {
        "score": sum(check["earned"] for check in checks),
        "checks": checks,
        "page_count": parsed.page_count,
        "columns": 2 if parsed.likely_multi_column else 1,
        "detected_sections": detected,
        "not_detected": [
            name for name in ("experience", "education", "skills") if name not in detected_set
        ],
        "warnings": warnings,
        "disclaimer": "CareerCanvas ATS Readiness Score is an application-specific heuristic. It is not a score produced by an employer's ATS.",
    }


def _merge(existing: dict, incoming: dict) -> dict:
    result = dict(existing)
    for key, value in incoming.items():
        current = result.get(key)
        if isinstance(value, list):
            combined = list(current or [])
            seen = {str(entry).lower() for entry in combined}
            for entry in value:
                marker = str(entry).lower()
                if marker not in seen:
                    combined.append(entry)
                    seen.add(marker)
            result[key] = combined
        elif not current and value:
            result[key] = value
    return result


@router.post("/preview")
def preview(payload: UploadPayload, db: Session = Depends(session)):
    parsed = parse_document(payload.filename, _safe_decode(payload))
    if not parsed.text.strip() and not parsed.scanned:
        raise HTTPException(422, "No readable text was found in the uploaded document")
    preamble, sections, detected = split_sections(parsed.text)
    personal = extract_personal(preamble, parsed.text)
    summary_lines = sections.get("summary", [])
    if summary_lines:
        personal["summary"] = {
            "value": " ".join(summary_lines)[:10000],
            "confidence": 0.88,
            "source": " | ".join(summary_lines)[:1000],
        }
    items = extract_items(sections)
    attach_duplicates(items, db)
    profile = get_profile(db)
    for key, candidate in personal.items():
        existing = str(profile.personal.get(key, ""))
        candidate["existing"] = existing
        candidate["conflict"] = bool(existing and _norm(existing) != _norm(candidate["value"]))
    return {
        "filename": payload.filename,
        "document": {
            "format": parsed.format,
            "page_count": parsed.page_count,
            "text_characters": len(parsed.text),
            "image_count": parsed.image_count,
            "link_count": parsed.link_count,
            "links": parsed.links,
            "table_count": parsed.table_count,
            "table_text_ratio": parsed.table_text_ratio,
            "scanned": parsed.scanned,
            "likely_multi_column": parsed.likely_multi_column,
        },
        "personal": personal,
        "items": items,
        "detected_sections": detected,
        "warnings": parsed.warnings,
        "extracted_text": parsed.text,
        "ats": uploaded_ats(parsed, detected, personal),
    }


@router.post("/apply")
def apply_import(payload: ApplyImport, db: Session = Depends(session)):
    profile = get_profile(db)
    personal = dict(profile.personal)
    personal_added = personal_updated = 0
    for key, selected in payload.personal_selected.items():
        if not selected or key not in PersonalDetails.model_fields:
            continue
        value = str(payload.personal.get(key, ""))[:10000]
        if not value:
            continue
        if not personal.get(key):
            personal_added += 1
        elif personal.get(key) != value:
            personal_updated += 1
        personal[key] = value
    profile.personal = PersonalDetails.model_validate(personal).model_dump()

    added = updated = skipped = 0
    selected_ids: list[str] = []
    for selection in payload.items:
        if not selection.selected or selection.kind not in KINDS:
            skipped += 1
            continue
        data = ItemInput(kind=selection.kind, data=selection.data).data
        duplicate = db.get(CareerItem, selection.duplicate_id) if selection.duplicate_id else None
        if duplicate is not None and duplicate.profile_id != profile.id:
            duplicate = None
        if duplicate is not None and selection.duplicate_action == "keep":
            selected_ids.append(duplicate.id)
            skipped += 1
            continue
        if duplicate is not None and selection.duplicate_action in {"replace", "merge"}:
            duplicate.data = data if selection.duplicate_action == "replace" else _merge(duplicate.data, data)
            duplicate.data = ItemInput(kind=duplicate.kind, data=duplicate.data).data
            selected_ids.append(duplicate.id)
            updated += 1
            continue
        item = CareerItem(profile_id=profile.id, kind=selection.kind, data=data)
        db.add(item)
        db.flush()
        selected_ids.append(item.id)
        added += 1

    db.add(
        AuditEvent(
            action=(
                f"Resume Imported {Path(payload.filename).name[:32]}: "
                f"{added} added, {updated} updated, {skipped} skipped"
            )[:100],
            entity_id=profile.id,
        )
    )
    db.commit()
    return {
        "personal_added": personal_added,
        "personal_updated": personal_updated,
        "items_added": added,
        "items_updated": updated,
        "items_skipped": skipped,
        "selected_ids": selected_ids,
    }
