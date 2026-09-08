from io import BytesIO
from urllib.parse import urlparse

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def hyperlink(paragraph, label, url, color='245c4b'):
    if urlparse(url).scheme not in {'http','https','mailto'}:
        return
    relationship=paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),relationship)
    run=OxmlElement('w:r');properties=OxmlElement('w:rPr');shade=OxmlElement('w:color');shade.set(qn('w:val'),color)
    properties.append(shade);run.append(properties);text=OxmlElement('w:t');text.text=label;run.append(text);link.append(run);paragraph._p.append(link)


def render_docx(data:dict):
    document=Document()
    personal=data['personal'];style=data['style'];hidden=set(personal.get('hidden_fields',[]))
    section=document.sections[0]
    section.page_width=Mm(210 if style['page_size']=='A4' else 215.9)
    section.page_height=Mm(297 if style['page_size']=='A4' else 279.4)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Mm(style['margin'])
    normal=document.styles['Normal'];normal.font.name='Arial' if style['font']=='System Sans' else style['font'];normal.font.size=Pt(style['font_size']);normal.font.color.rgb=RGBColor.from_string(style['text'][1:])
    normal.paragraph_format.line_spacing=style['line_height'];normal.paragraph_format.space_after=Pt(4)
    normal.paragraph_format.widow_control=True
    for name in ['Title','Subtitle','Heading 1','Heading 2']:
        document.styles[name].font.name=normal.font.name
        document.styles[name].font.color.rgb=RGBColor.from_string(style['accent'][1:])
    document.styles['Title'].font.size=Pt(27)
    document.styles['Title'].font.color.rgb=RGBColor(0,0,0)
    for border in document.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    document.styles['Heading 1'].font.size=Pt(style['font_size']+1)
    document.styles['Heading 1'].paragraph_format.space_before=Pt(style['section_spacing']*.75)
    document.styles['Heading 1'].paragraph_format.space_after=Pt(6)
    document.styles['Heading 1'].paragraph_format.keep_with_next=True
    title=document.add_paragraph(personal.get('full_name',''),style='Title')
    if style['alignment']=='center':title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if 'professional_title' not in hidden and personal.get('professional_title'):
        document.add_paragraph(personal['professional_title'],style='Subtitle')
    contact=document.add_paragraph()
    if personal.get('email') and 'email' not in hidden:hyperlink(contact,personal['email'],'mailto:'+personal['email'],style['accent'][1:])
    for key in ['phone','city','country']:
        if personal.get(key) and key not in hidden:contact.add_run('  |  '+personal[key])
    for key in ['linkedin','github','portfolio','website']:
        if personal.get(key) and key not in hidden:
            contact.add_run('  |  ');hyperlink(contact,key.title(),personal[key],style['accent'][1:])
    if personal.get('summary') and 'summary' not in hidden:
        document.add_heading('Professional Summary',level=1)
        if data.get('summary_rich'):rich_docx(document,data['summary_rich'])
        else:document.add_paragraph(personal['summary'])
    for sec in data['sections']:
        if not sec['visible']:continue
        document.add_heading(sec['heading'],level=1)
        if sec['kind']=='skills':
            document.add_paragraph(sec.get('separator',' · ').join(item['data']['name'] for item in sec['items']));continue
        for item in sec['items']:
            item_start=len(document.paragraphs)
            d=item['data']
            heading=d.get('position') or d.get('name') or d.get('title') or d.get('statement') or d.get('language') or d.get('institution') or ''
            p=document.add_paragraph();p.add_run(heading).bold=True;p.paragraph_format.keep_with_next=True
            start=d.get('start_date') or d.get('date') or d.get('year') or ''
            end='Present' if d.get('current') else d.get('end_date','')
            if start or end:p.add_run('  |  '+start+(' - '+end if end else '')).italic=True
            details=' · '.join(str(d[k]) for k in ['company','issuer','venue','role','location'] if d.get(k))
            if details:document.add_paragraph(details)
            if d.get('degree'):document.add_paragraph(' · '.join(str(d[k]) for k in ['degree','field','gpa'] if d.get(k)))
            for key in ['authors','description','citation','proficiency']:
                if d.get(key):document.add_paragraph(str(d[key]))
            for bullet in d.get('bullets',[]):
                p=document.add_paragraph(bullet['text'],style='List Bullet');p.paragraph_format.space_after=Pt(style['bullet_spacing']*.75)
            for key in ['technologies','courses','honors']:
                if d.get(key):document.add_paragraph(key.title()+': '+', '.join(d[key]))
            for key in ['url','github','credential_url']:
                if d.get(key):hyperlink(document.add_paragraph(),d[key],d[key],style['accent'][1:])
            if d.get('doi'):document.add_paragraph('DOI: '+d['doi'])
            if d.get('credential_id'):document.add_paragraph('Credential ID: '+d['credential_id'])
            if sec['kind']=='references':
                if d.get('email'):hyperlink(document.add_paragraph(),d['email'],'mailto:'+d['email'])
                if d.get('phone'):document.add_paragraph(d['phone'])
            # Keep a normal-sized entry together, including its final hyperlinks.
            item_paragraphs=document.paragraphs[item_start:]
            for paragraph in item_paragraphs[:-1]:paragraph.paragraph_format.keep_with_next=True
            if item_paragraphs:item_paragraphs[-1].paragraph_format.keep_with_next=False
    document.core_properties.author=''
    document.core_properties.last_modified_by=''
    document.core_properties.title=personal.get('full_name','')+' Resume'
    document.core_properties.subject=''
    document.core_properties.comments=''
    output=BytesIO();document.save(output);return output.getvalue()


def rich_docx(document,node,paragraph=None,bullet=False):
    kind=node['type']
    if kind=='text':
        if paragraph is None:paragraph=document.add_paragraph()
        run=paragraph.add_run(node.get('text') or '')
        for mark in node.get('marks') or []:
            if mark['type']=='bold':run.bold=True
            if mark['type']=='italic':run.italic=True
    elif kind=='hardBreak' and paragraph is not None:paragraph.add_run().add_break()
    else:
        if kind=='paragraph':paragraph=document.add_paragraph(style='List Bullet' if bullet else 'Normal')
        for child in node.get('content') or []:rich_docx(document,child,paragraph,bullet or kind=='bulletList')
