from backend.app.resume_schemas import ResumeDocument
from backend.app.docx_export import render_docx
from docx import Document
from io import BytesIO
import pytest


def test_rich_summary_safe_and_docx():
    rich = {
        "type": "doc",
        "content": [
            {
                "type": "paragraph",
                "content": [
                    {
                        "type": "text",
                        "text": "Built reliable systems",
                        "marks": [{"type": "bold"}],
                    }
                ],
            }
        ],
    }
    doc = ResumeDocument(
        personal={"full_name": "Alex", "summary": "Built reliable systems"},
        summary_rich=rich,
    ).model_dump()
    parsed = Document(BytesIO(render_docx(doc)))
    assert any(
        r.bold and r.text == "Built reliable systems"
        for p in parsed.paragraphs
        for r in p.runs
    )
    with pytest.raises(ValueError):
        ResumeDocument(summary_rich={"type": "script", "text": "unsafe"})
