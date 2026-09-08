from io import BytesIO
from zipfile import ZipFile

from docx import Document
import pytest

from backend.app.config import ROOT
from backend.app.resume_schemas import TEMPLATES
from backend.tests.test_resumes import make_resume


@pytest.mark.parametrize('template',TEMPLATES)
def test_docx_is_editable_and_preserves_sections_bullets_and_links(client,template):
    resume=make_resume(client)
    resume['document']['template']=template
    assert client.put(f'/api/resumes/{resume["id"]}',json={key:resume[key] for key in ['name','revision','target_role','archived','document']}).status_code==200
    response=client.post(f'/api/resumes/{resume["id"]}/export/docx')
    assert response.status_code==200
    document=Document(BytesIO(response.content))
    text='\n'.join(p.text for p in document.paragraphs)
    for expected in ['Alex Morgan','Experience','Education','Projects','Skills','Python','Northstar Labs','1,200','Machine Learning']:
        assert expected in text
    assert any(p.style.name=='List Bullet' for p in document.paragraphs)
    assert any(p.style.name=='Heading 1' for p in document.paragraphs)
    assert abs(document.sections[0].page_width.mm-210)<.1
    with ZipFile(BytesIO(response.content)) as archive:
        relationships=archive.read('word/_rels/document.xml.rels').decode()
        assert 'https://example.com/signal' in relationships
        assert 'mailto:alex.morgan@example.com' in relationships
        assert not any(name.startswith('word/media/') for name in archive.namelist())
    if template=='Modern':
        target=ROOT/'data/validation/docx';target.mkdir(parents=True,exist_ok=True);(target/'demo-resume.docx').write_bytes(response.content)


def test_hidden_content_not_exported(client):
    resume=make_resume(client)
    resume['document']['personal']['hidden_fields']=['email','phone']
    for sec in resume['document']['sections']:
        if sec['kind']=='certifications':sec['visible']=False
    from backend.app.docx_export import render_docx
    result=render_docx(resume['document'])
    doc=Document(BytesIO(result))
    assert 'Applied Machine Learning Certificate' not in '\n'.join(p.text for p in doc.paragraphs)
    with ZipFile(BytesIO(result)) as archive:assert 'mailto:alex.morgan@example.com' not in archive.read('word/_rels/document.xml.rels').decode()
