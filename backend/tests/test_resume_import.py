import base64
import io

from docx import Document
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from backend.app.resume_import import parse_document, split_sections


STANDARD_TEXT = """Alex Morgan
AI Engineer
alex.morgan@example.com
+1 555 123 4567
https://github.com/alexmorgan
Seattle, USA

Professional Summary
AI engineer focused on reliable production systems.

Experience
AI Engineer | Acme Systems
Jan 2024 - Present
Built FastAPI services used by three internal products.
Reduced model evaluation time by 35 percent.

Education
Example University
MS Artificial Intelligence
2022 - 2024
GPA: 3.9/4.0

Technical Skills
Python, FastAPI, SQL, Docker

Projects
CareerCanvas
Built a local-first resume and career workspace.
"""


def encoded(text: str) -> str:
    return base64.b64encode(text.encode()).decode()


def encoded_bytes(data: bytes) -> str:
    return base64.b64encode(data).decode()


def text_pdf(text: str) -> bytes:
    """Create a deterministic selectable-text PDF without a test-only PDF dependency."""
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
    )
    operations = ["BT", "/F1 10 Tf", "72 740 Td"]
    for index, line in enumerate(text.splitlines()):
        if index:
            operations.append("0 -14 Td")
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        operations.append(f"({escaped}) Tj")
    operations.append("ET")
    content = DecodedStreamObject()
    content.set_data("\n".join(operations).encode("latin-1"))
    page[NameObject("/Contents")] = writer._add_object(content)
    stream = io.BytesIO()
    writer.write(stream)
    assert "Alex Morgan" in PdfReader(io.BytesIO(stream.getvalue())).pages[0].extract_text()
    return stream.getvalue()


def docx_bytes(text: str, *, use_table: bool = False) -> bytes:
    document = Document()
    if use_table:
        lines = [line for line in text.splitlines() if line.strip()]
        table = document.add_table(rows=len(lines), cols=1)
        for index, line in enumerate(lines):
            table.cell(index, 0).text = line
    else:
        for line in text.splitlines():
            if line.strip():
                document.add_paragraph(line)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_txt_preview_extracts_profile_and_direct_ats(client):
    response = client.post(
        "/api/import/preview",
        json={"filename": "standard_resume.txt", "content_base64": encoded(STANDARD_TEXT)},
    )
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["personal"]["email"]["value"] == "alex.morgan@example.com"
    assert data["personal"]["phone"]["value"]
    assert {"experience", "education", "skills", "projects"}.issubset(
        set(data["detected_sections"])
    )
    assert any(
        item["kind"] == "skills" and item["data"]["name"] == "Python"
        for item in data["items"]
    )
    assert data["ats"]["score"] >= 80
    assert not data["ats"]["not_detected"]
    assert "AI engineer focused" in data["extracted_text"]


def test_selectable_pdf_runs_direct_ats_and_preserves_detected_link(client):
    response = client.post(
        "/api/import/preview",
        json={
            "filename": "standard_resume.pdf",
            "content_base64": encoded_bytes(text_pdf(STANDARD_TEXT)),
        },
    )
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["document"]["format"] == "PDF"
    assert data["document"]["page_count"] == 1
    assert data["document"]["scanned"] is False
    assert "https://github.com/alexmorgan" in data["document"]["links"]
    assert data["personal"]["email"]["value"] == "alex.morgan@example.com"
    assert data["ats"]["score"] >= 80
    assert "Experience" in data["extracted_text"]


def test_docx_upload_runs_direct_ats(client):
    response = client.post(
        "/api/import/preview",
        json={
            "filename": "minimal_resume.docx",
            "content_base64": encoded_bytes(docx_bytes(STANDARD_TEXT)),
        },
    )
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["document"]["format"] == "DOCX"
    assert {"experience", "education", "skills"}.issubset(set(data["detected_sections"]))
    assert data["personal"]["email"]["value"] == "alex.morgan@example.com"
    assert data["ats"]["score"] >= 80


def test_import_review_applies_then_detects_duplicates(client):
    preview = client.post(
        "/api/import/preview",
        json={"filename": "standard_resume.txt", "content_base64": encoded(STANDARD_TEXT)},
    ).json()
    payload = {
        "filename": preview["filename"],
        "personal": {key: value["value"] for key, value in preview["personal"].items()},
        "personal_selected": {key: True for key in preview["personal"]},
        "items": [
            {
                "temp_id": item["temp_id"],
                "kind": item["kind"],
                "data": item["data"],
                "selected": True,
                "duplicate_id": None,
                "duplicate_action": "new",
            }
            for item in preview["items"]
        ],
    }
    applied = client.post("/api/import/apply", json=payload)
    assert applied.status_code == 200, applied.text
    result = applied.json()
    assert result["items_added"] >= 4
    profile = client.get("/api/profile").json()
    assert profile["personal"]["full_name"] == "Alex Morgan"
    assert any(
        item["kind"] == "skills" and item["data"]["name"] == "Python"
        for item in profile["items"]
    )

    second = client.post(
        "/api/import/preview",
        json={"filename": "standard_resume.txt", "content_base64": encoded(STANDARD_TEXT)},
    ).json()
    python = next(
        item
        for item in second["items"]
        if item["kind"] == "skills" and item["data"]["name"] == "Python"
    )
    assert python["duplicate"]
    assert python["duplicate"]["score"] == 1.0


def test_duplicate_keep_does_not_create_second_skill(client):
    first = client.post(
        "/api/import/preview",
        json={"filename": "standard_resume.txt", "content_base64": encoded(STANDARD_TEXT)},
    ).json()
    item = next(
        x for x in first["items"] if x["kind"] == "skills" and x["data"]["name"] == "Python"
    )
    client.post(
        "/api/import/apply",
        json={
            "filename": first["filename"],
            "personal": {},
            "personal_selected": {},
            "items": [
                {
                    "temp_id": item["temp_id"],
                    "kind": item["kind"],
                    "data": item["data"],
                    "selected": True,
                    "duplicate_id": None,
                    "duplicate_action": "new",
                }
            ],
        },
    )
    second = client.post(
        "/api/import/preview",
        json={"filename": "standard_resume.txt", "content_base64": encoded(STANDARD_TEXT)},
    ).json()
    duplicate = next(
        x for x in second["items"] if x["kind"] == "skills" and x["data"]["name"] == "Python"
    )
    response = client.post(
        "/api/import/apply",
        json={
            "filename": second["filename"],
            "personal": {},
            "personal_selected": {},
            "items": [
                {
                    "temp_id": duplicate["temp_id"],
                    "kind": duplicate["kind"],
                    "data": duplicate["data"],
                    "selected": True,
                    "duplicate_id": duplicate["duplicate"]["id"],
                    "duplicate_action": "keep",
                }
            ],
        },
    )
    assert response.status_code == 200
    assert response.json()["items_skipped"] == 1
    skills = [
        x
        for x in client.get("/api/profile").json()["items"]
        if x["kind"] == "skills" and x["data"]["name"] == "Python"
    ]
    assert len(skills) == 1


def test_missing_phone_is_reported_by_uploaded_ats(client):
    text = STANDARD_TEXT.replace("+1 555 123 4567\n", "")
    data = client.post(
        "/api/import/preview",
        json={"filename": "missing_phone_resume.txt", "content_base64": encoded(text)},
    ).json()
    phone = next(check for check in data["ats"]["checks"] if check["key"] == "phone")
    assert phone["passed"] is False
    assert data["ats"]["score"] < 100


def test_year_range_is_not_imported_as_phone(client):
    text = STANDARD_TEXT.replace("+1 555 123 4567\n", "")
    data = client.post(
        "/api/import/preview",
        json={"filename": "year_ranges_only.txt", "content_base64": encoded(text)},
    ).json()
    assert "phone" not in data["personal"]


def test_docx_extraction_and_table_warning():
    parsed = parse_document("resume.docx", docx_bytes(STANDARD_TEXT, use_table=True))
    assert parsed.format == "DOCX"
    assert "Alex Morgan" in parsed.text
    assert parsed.table_count == 1
    assert parsed.table_text_ratio > 0
    assert parsed.likely_multi_column is True
    assert parsed.warnings


def test_scanned_pdf_is_detected_without_cloud_ocr():
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    stream = io.BytesIO()
    writer.write(stream)
    parsed = parse_document("image_heavy_resume.pdf", stream.getvalue())
    assert parsed.scanned is True
    assert any("OCR" in warning for warning in parsed.warnings)


def test_heading_aliases_are_normalized():
    _, sections, detected = split_sections(
        "Alex Morgan\nPROFESSIONAL EXPERIENCE\nEngineer | Acme\n2023 - 2024\nTECHNICAL SKILLS\nPython"
    )
    assert "experience" in detected
    assert "skills" in detected
    assert sections["skills"] == ["Python"]
